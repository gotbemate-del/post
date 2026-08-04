# -*- coding: utf-8 -*-
"""把原始 Excel 名單解析成以「街道」分類的結構化資料 (data/stores.json)。

原始檔規則：
  * 分頁 旅行社 / 旅店 / 餐廳 為澎湖店家名單，欄位順序各不相同。
  * 整列有底色 = 已經發送過的店家（實際用到的是黃色 FFFF00 與米色 FFF2CC；
    表頭／空白列的灰底 FFD9D9D9 與白底不算，見 IGNORED_FILLS）。
  * 分頁「不再名單上新增的」為手動補充、尚未整理進名單的店家。

用法：
    python3 scripts/parse_xlsx.py [來源.xlsx] [輸出.json]
"""
import json
import re
import sys
from collections import Counter, OrderedDict, defaultdict

import openpyxl

DEFAULT_SRC = 'data/原始名單.xlsx'
DEFAULT_OUT = 'data/stores.json'

# 表頭與空白列使用的灰底，不代表「已發送」
IGNORED_FILLS = {None, '00000000', 'FFFFFFFF', 'FFD9D9D9'}

# 非路街地名的正規化：同一地點的不同寫法、舊地名、涵蓋範圍較小的聚落名
AREA_ALIAS = {
    '西文澳': '西文', '東文澳': '東文',
    '大案山': '案山',
    '鎖管港': '鎖港',
    '大赤崁': '赤崁',
    '鐵線尾': '鐵線',
    '港底': '成功', '土地公前': '沙港',
    '南滬': '南港',
}

NO_STREET = '地址未含街道'

# 優先跑的市區路段，固定排在所有街道最前面，順序就照這裡列的
PRIORITY_STREETS = [
    ('馬公市', '中正路'),
    ('馬公市', '民權路'),
    ('馬公市', '中華路'),
    ('馬公市', '陽明路'),
    ('馬公市', '三多路'),
]

# 各分頁的欄位對應（欄位順序在每張分頁都不一樣，不能共用）
SHEETS = OrderedDict([
    ('旅行社', OrderedDict([('項次', 1), ('項目', 2), ('名稱', 3), ('地址', 4),
                            ('電話', 5), ('營業時間', 6), ('連結', 7)])),
    ('旅店', OrderedDict([('項次', 1), ('項目', 2), ('名稱', 3), ('地址', 4),
                          ('電話', 5), ('備註', 6), ('照片', 7), ('連結', 8)])),
    ('餐廳', OrderedDict([('項次', 1), ('項目', 2), ('名稱', 3), ('地址', 4),
                          ('電話', 5), ('備註', 6), ('營業時間', 7), ('連結', 8)])),
])

EXTRA_SHEET = '不再名單上新增的'


def clean(value):
    if value is None:
        return ''
    s = str(value).strip()
    return '' if s.lower() in ('none', 'nan') else s


def is_sent(cell):
    """整列有黃 / 米底色即代表已發送過。"""
    f = cell.fill
    if not f or f.fill_type != 'solid' or not f.fgColor:
        return False
    return f.fgColor.rgb not in IGNORED_FILLS


def parse_address(raw):
    """把地址拆成 (鄉鎮市, 街道或地區)。

    澎湖地址有兩種寫法：有路名的（馬公市三多路353號）與只有村里聚落的
    （馬公市西衛里223之1號）。兩種都要能歸到同一個「街道」維度。
    """
    if not raw:
        return '', NO_STREET
    body = clean(raw).replace(' ', '').replace('　', '')
    body = re.sub(r'^\d{3,6}', '', body)                       # 去郵遞區號
    body = body.replace('台湾', '').replace('台灣', '').replace('臺灣', '')
    body = re.sub(r'^[臺台]?澎湖縣?', '', body)

    # 取最內層的鄉鎮市：少數地址寫成「馬公市白沙鄉赤崁村…」，後者才是實際位置
    town = ''
    while True:
        m = re.match(r'^([一-鿿]{2,3}[市鄉鎮])', body)
        if not m:
            break
        town = m.group(1)
        body = body[m.end():]

    # 里名後面還接著路名時（西衛里西衛路），以路名為準
    m = re.match(r'^([一-鿿]{1,4}里)(?=[一-鿿])', body)
    lead_village = ''
    if m:
        lead_village = m.group(1)
        rest = body[m.end():]
        if re.match(r'^[一-鿿]{1,6}?(?:大道|路|街)', rest):
            body = rest

    m = re.match(r'^([一-鿿]{1,6}?(?:大道|路|街))((?:[一二三四五六七八九十]+段)?)', body)
    if m:
        return town, m.group(1) + m.group(2)

    # 沒有路名 → 用村里 / 聚落名
    m = re.match(r'^([一-鿿]+?)([里村])([一-鿿]*)', body)
    if m:
        base, suffix, sub = m.group(1), m.group(2), m.group(3)
        sub = re.sub(r'(之|號|樓|巷|弄).*$', '', sub)
        if sub:
            name = AREA_ALIAS.get(sub, base)
        elif len(base) >= 3:
            name = base + suffix        # 漁港新村這類三字以上的完整村名
        else:
            name = base
        return town, AREA_ALIAS.get(name, name)

    m = re.match(r'^([一-鿿]+)', body)
    if m:
        name = re.sub(r'(之|號|樓|巷|弄).*$', '', m.group(1))
        if name:
            return town, AREA_ALIAS.get(name, name)

    if lead_village:
        return town, lead_village.rstrip('里')
    return town, NO_STREET


def dedupe_key(name):
    """判斷是否為同一家店：去掉公司型態與標點後比對名稱。"""
    key = re.sub(r'[\s()（）\-·．.,，]', '', name)
    key = re.sub(r'(股份)?有限公司', '', key)
    key = re.sub(r'(澎湖)?分公司', '', key)
    return key


def read_records(wb):
    records = []
    for sheet, cols in SHEETS.items():
        ws = wb[sheet]
        for row in range(2, ws.max_row + 1):
            name = clean(ws.cell(row, cols['名稱']).value)
            if not name:
                continue
            get = lambda k: clean(ws.cell(row, cols[k]).value) if k in cols else ''
            town, street = parse_address(ws.cell(row, cols['地址']).value)
            records.append({
                'id': f'{sheet}-{row}',
                'sheet': sheet,
                'row': row,
                'category': get('項目') or sheet,
                'name': name,
                'address': clean(ws.cell(row, cols['地址']).value),
                'town': town or '其他',
                'street': street,
                'phone': get('電話'),
                'hours': get('營業時間'),
                'note': get('備註'),
                'link': get('連結'),
                'sent': any(is_sent(ws.cell(row, c)) for c in range(1, len(cols) + 1)),
            })
    return records


def link_duplicates(records):
    """同一家店在多張名單重複出現時互相連結，並讓「已發送」狀態同步。"""
    groups = defaultdict(list)
    for rec in records:
        groups[dedupe_key(rec['name'])].append(rec)
    for key, group in groups.items():
        if len(group) < 2:
            continue
        sent = any(r['sent'] for r in group)
        for rec in group:
            rec['sent'] = sent
            rec['duplicate_of'] = [r['id'] for r in group if r['id'] != rec['id']]


def read_extras(wb):
    ws = wb[EXTRA_SHEET]
    extras = []
    for col in (2, 4):
        batch = clean(ws.cell(2, col).value)
        for row in range(3, ws.max_row + 1):
            text = clean(ws.cell(row, col).value)
            if text:
                extras.append({'batch': batch, 'text': re.sub(r'^\d+\.\s*', '', text)})
    return extras


def build_streets(records):
    """依鄉鎮市 → 街道分組，店家多的街道排前面。"""
    buckets = defaultdict(list)
    for rec in records:
        buckets[(rec['town'], rec['street'])].append(rec)

    def sort_key(item):
        (town, street), items = item
        # 優先路段照 PRIORITY_STREETS 的順序排在最前，其餘沿用原本的規則
        rank = (PRIORITY_STREETS.index((town, street))
                if (town, street) in PRIORITY_STREETS else len(PRIORITY_STREETS))
        return (rank, town != '馬公市', town, -len(items), street)

    streets = []
    for (town, street), items in sorted(buckets.items(), key=sort_key):
        items.sort(key=lambda r: (r['category'], r['name']))
        streets.append({
            'town': town,
            'street': street,
            'key': f'{town}|{street}',
            'total': len(items),
            'sent': sum(1 for r in items if r['sent']),
            'store_ids': [r['id'] for r in items],
        })
    return streets


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SRC
    out = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUT

    wb = openpyxl.load_workbook(src)
    records = read_records(wb)
    link_duplicates(records)
    streets = build_streets(records)
    extras = read_extras(wb)

    payload = {
        'source': src,
        'generated_by': 'scripts/parse_xlsx.py',
        'summary': {
            'stores': len(records),
            'sent': sum(1 for r in records if r['sent']),
            'streets': len(streets),
            'towns': len({s['town'] for s in streets}),
            'by_category': dict(Counter(r['category'] for r in records)),
        },
        'streets': streets,
        'stores': records,
        'extras': extras,
    }
    with open(out, 'w', encoding='utf-8') as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=1)

    s = payload['summary']
    print(f"已輸出 {out}：{s['stores']} 家店、{s['sent']} 家已發送、"
          f"{s['streets']} 條街道、{s['towns']} 個鄉鎮市")


if __name__ == '__main__':
    main()
