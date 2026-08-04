# -*- coding: utf-8 -*-
"""由 data/stores.json 產生「依街道分類」的確認表單 Excel。

輸出三張分頁：
  1. 街道總覽 — 每條街道的店家數、已發送數、發送率
  2. 確認表單 — 主表，依街道分組列出所有店家與待填欄位
  3. 待補名單 — 原始檔「不再名單上新增的」分頁的手動補充項目

用法：
    python3 scripts/build_form_xlsx.py [輸入.json] [輸出.xlsx]
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DEFAULT_IN = 'data/stores.json'
DEFAULT_OUT = '澎湖店家確認表單.xlsx'

FONT = 'Arial'
HEADER_FILL = PatternFill('solid', fgColor='FF1F3864')
STREET_FILL = PatternFill('solid', fgColor='FFDDEBF7')
SENT_FILL = PatternFill('solid', fgColor='FFFFF2CC')       # 沿用原始檔的「已發送」底色
INPUT_FILL = PatternFill('solid', fgColor='FFFFFF00')      # 待填欄位
THIN = Side(style='thin', color='FFBFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_OPTIONS = ['未發送', '已發送']
CHANNEL_OPTIONS = ['', '現場拜訪', '電話', 'LINE', 'Email', 'FB/IG 私訊']

FORM_HEADERS = [
    ('街道', 14), ('鄉鎮市', 10), ('分類', 12), ('店家名稱', 34), ('地址', 40),
    ('電話', 16), ('營業時間', 26), ('原始狀態', 10),
    ('發送狀態', 12), ('發送日期', 12), ('發送方式', 12), ('負責人', 10),
    ('回覆／備註', 30), ('連結', 34),
]


def style_header(ws, headers, row=1):
    for idx, (title, width) in enumerate(headers, start=1):
        cell = ws.cell(row, idx, title)
        cell.font = Font(name=FONT, bold=True, color='FFFFFFFF', size=11)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[row].height = 26


def build_overview(wb, data):
    ws = wb.create_sheet('街道總覽')
    headers = [('鄉鎮市', 12), ('街道／地區', 18), ('店家數', 10),
               ('已發送', 10), ('未發送', 10), ('發送率', 12)]

    ws['A1'] = '澎湖店家確認表單 — 街道總覽'
    ws['A1'].font = Font(name=FONT, bold=True, size=14)
    s = data['summary']
    ws['A2'] = (f"共 {s['stores']} 家店家，分布於 {s['towns']} 個鄉鎮市、{s['streets']} 條街道／地區；"
                f"原始檔標記為已發送者 {s['sent']} 家。")
    ws['A2'].font = Font(name=FONT, size=10, color='FF595959')
    ws['A3'] = '「原始狀態」欄的「已發送」來自原始 Excel 中有底色的列。'
    ws['A3'].font = Font(name=FONT, size=10, color='FF595959')

    style_header(ws, headers, row=5)
    row = 6
    for st in data['streets']:
        ws.cell(row, 1, st['town'])
        ws.cell(row, 2, st['street'])
        ws.cell(row, 3, st['total'])
        ws.cell(row, 4, st['sent'])
        ws.cell(row, 5, f'=C{row}-D{row}')
        ws.cell(row, 6, f'=IF(C{row}=0,0,D{row}/C{row})')
        ws.cell(row, 6).number_format = '0.0%'
        for col in range(1, 7):
            cell = ws.cell(row, col)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='center' if col >= 3 else 'left',
                                       vertical='center')
        row += 1

    last = row - 1
    ws.cell(row, 2, '合計').font = Font(name=FONT, bold=True)
    ws.cell(row, 3, f'=SUM(C6:C{last})').font = Font(name=FONT, bold=True)
    ws.cell(row, 4, f'=SUM(D6:D{last})').font = Font(name=FONT, bold=True)
    ws.cell(row, 5, f'=SUM(E6:E{last})').font = Font(name=FONT, bold=True)
    ws.cell(row, 6, f'=IF(C{row}=0,0,D{row}/C{row})')
    ws.cell(row, 6).font = Font(name=FONT, bold=True)
    ws.cell(row, 6).number_format = '0.0%'
    for col in range(1, 7):
        ws.cell(row, col).border = BORDER

    ws.freeze_panes = 'A6'
    ws.auto_filter.ref = f'A5:F{last}'
    return ws


def build_form(wb, data):
    ws = wb.create_sheet('確認表單')
    by_id = {s['id']: s for s in data['stores']}

    ws['A1'] = '澎湖店家確認表單（依街道分類）'
    ws['A1'].font = Font(name=FONT, bold=True, size=14)
    ws['A2'] = ('填寫說明：黃底欄位（I～M 欄）為待填欄位，其餘為原始資料請勿修改。'
                '「原始狀態」為已發送者整列以米色標示。')
    ws['A2'].font = Font(name=FONT, size=10, color='FF595959')

    header_row = 4
    style_header(ws, FORM_HEADERS, row=header_row)

    row = header_row + 1
    data_rows = []
    for st in data['streets']:
        # 街道分隔列
        label = f"{st['town']} ── {st['street']}（{st['total']} 家，已發送 {st['sent']} 家）"
        ws.cell(row, 1, label)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(FORM_HEADERS))
        cell = ws.cell(row, 1)
        cell.font = Font(name=FONT, bold=True, size=11, color='FF1F3864')
        cell.fill = STREET_FILL
        cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        for col in range(1, len(FORM_HEADERS) + 1):
            ws.cell(row, col).border = BORDER
        row += 1

        for sid in st['store_ids']:
            store = by_id[sid]
            note = store['note']
            if store.get('duplicate_of'):
                dup_names = '、'.join(by_id[d]['sheet'] for d in store['duplicate_of'])
                note = (note + ' / ' if note else '') + f'※ 與「{dup_names}」名單重複'
            values = [
                st['street'], st['town'], store['category'], store['name'], store['address'],
                store['phone'], store['hours'], '已發送' if store['sent'] else '未發送',
                '已發送' if store['sent'] else '未發送', None, None, None, note, store['link'],
            ]
            for idx, value in enumerate(values, start=1):
                cell = ws.cell(row, idx, value)
                cell.font = Font(name=FONT, size=10)
                cell.border = BORDER
                cell.alignment = Alignment(vertical='center', wrap_text=idx in (5, 7, 13))
            if store['sent']:
                for col in range(1, len(FORM_HEADERS) + 1):
                    ws.cell(row, col).fill = SENT_FILL
            for col in range(9, 14):                      # I～M 待填欄位
                ws.cell(row, col).fill = INPUT_FILL
            ws.cell(row, 10).number_format = 'yyyy/mm/dd'
            data_rows.append(row)
            row += 1

    last = row - 1

    # 街道分隔列不需要下拉選單，只套用在實際的店家列上
    def sqref(column):
        parts, start, prev = [], None, None
        for r in data_rows:
            if start is None:
                start, prev = r, r
            elif r == prev + 1:
                prev = r
            else:
                parts.append(f'{column}{start}:{column}{prev}')
                start, prev = r, r
        if start is not None:
            parts.append(f'{column}{start}:{column}{prev}')
        return parts

    status_dv = DataValidation(type='list', formula1=f'"{",".join(STATUS_OPTIONS)}"',
                               allow_blank=True, showDropDown=False)
    channel_dv = DataValidation(type='list', formula1=f'"{",".join(CHANNEL_OPTIONS[1:])}"',
                                allow_blank=True, showDropDown=False)
    ws.add_data_validation(status_dv)
    ws.add_data_validation(channel_dv)
    status_ranges = sqref('I')
    for part in status_ranges:
        status_dv.add(part)
    for part in sqref('K'):
        channel_dv.add(part)

    for value, fill_color, font_color in (('已發送', 'FFC6EFCE', 'FF006100'),
                                          ('未發送', 'FFF2F2F2', 'FF808080')):
        ws.conditional_formatting.add(
            ' '.join(status_ranges),
            CellIsRule(operator='equal', formula=[f'"{value}"'],
                       fill=PatternFill('solid', fgColor=fill_color), font=Font(color=font_color)))

    ws.freeze_panes = f'A{header_row + 1}'
    ws.auto_filter.ref = f'A{header_row}:N{last}'
    return ws


def build_extras(wb, data):
    ws = wb.create_sheet('待補名單')
    ws['A1'] = '不在原始名單上、後續手動新增的店家'
    ws['A1'].font = Font(name=FONT, bold=True, size=14)
    ws['A2'] = '這些項目原始檔只有一行文字，沒有地址，因此無法歸入街道，待補齊地址後再併入確認表單。'
    ws['A2'].font = Font(name=FONT, size=10, color='FF595959')

    headers = [('批次', 26), ('店家（原始文字）', 34), ('地址（待補）', 40),
               ('電話（待補）', 18), ('發送狀態', 12), ('備註', 26)]
    style_header(ws, headers, row=4)

    row = 5
    for extra in data['extras']:
        ws.cell(row, 1, extra['batch'])
        ws.cell(row, 2, extra['text'])
        for col in range(1, 7):
            cell = ws.cell(row, col)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(vertical='center')
            if col >= 3:
                cell.fill = INPUT_FILL
        row += 1

    last = row - 1
    if last >= 5:
        dv = DataValidation(type='list', formula1=f'"{",".join(STATUS_OPTIONS)}"',
                            allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(f'E5:E{last}')
        ws.auto_filter.ref = f'A4:F{last}'
    ws.freeze_panes = 'A5'
    return ws


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_IN
    out = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUT

    with open(src, encoding='utf-8') as fh:
        data = json.load(fh)

    wb = Workbook()
    wb.remove(wb.active)
    build_overview(wb, data)
    build_form(wb, data)
    build_extras(wb, data)
    wb.save(out)
    print(f'已輸出 {out}')


if __name__ == '__main__':
    main()
